"""Export round-trip (verification 7): fake runs -> export -> recalculate -> 0 MISMATCH, numbers equal the UI preview."""
import sys
import tempfile
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import config  # noqa: E402
import exporter  # noqa: E402
import fake_data  # noqa: E402
import matching  # noqa: E402
import scoring  # noqa: E402

FORMULA_CELLS_LOG = ["I", "J", "N", "Q", "Z"]


@pytest.fixture(scope="module")
def exported():
    conn, _ = fake_data.build()
    out = Path(tempfile.mkdtemp(prefix="alpha_exports_"))
    sha_before = exporter.sha256_file(config.TEMPLATE)
    res = exporter.build_export(conn, reviewed_only=True, out_dir=out, stamp="TEST")
    return conn, res, sha_before


def test_outcome_coverage(exported):
    conn, res, _ = exported
    outcomes = {r["outcome"] for r in res["records"]}
    assert {"TP", "FN", "TN", "FP", "Duplicate", "Late", "Excluded", "Needs start"} <= outcomes
    assert any(r["alerted_instead"] == "Smoking" for r in res["records"])
    assert any(r["added_after_run"] for r in res["records"])
    assert any(r["set"] == "Validation" and r["outcome"] == "TP" for r in res["records"])


def test_template_unchanged(exported):
    _, res, sha_before = exported
    assert res["template_sha256"] == sha_before == exporter.sha256_file(config.TEMPLATE)


def test_formulas_preserved(exported):
    _, res, _ = exported
    wb = openpyxl.load_workbook(res["path"])
    tpl = openpyxl.load_workbook(config.TEMPLATE)
    ws, wt = wb["Results Log"], tpl["Results Log"]
    for r in range(5, 405):
        for col in FORMULA_CELLS_LOG:
            assert ws[f"{col}{r}"].value == wt[f"{col}{r}"].value
    for name in ("Summary", "Validation", "Breakdown", "Confusion Matrix"):
        for row in wb[name].iter_rows():
            for c in row:
                assert c.value == tpl[name][c.coordinate].value, (name, c.coordinate)
    for r in range(5, 155):
        assert wb["Clips"][f"N{r}"].value == tpl["Clips"][f"N{r}"].value
    assert wb.sheetnames == tpl.sheetnames


def test_recalculated_check_and_preview(exported):
    _, res, _ = exported
    if res["method"] is None:
        pytest.skip("no recalculation engine available")
    assert res["mismatches"] == [], res["mismatches"]
    assert res["preview_diffs"] == [], res["preview_diffs"][:5]


def test_values_survive_reload(exported):
    _, res, _ = exported
    wb = openpyxl.load_workbook(res["path"])
    ws = wb["Results Log"]
    first = [ws.cell(5, c).value for c in (1, 2, 3, 4, 7, 8, 11, 12, 25)]
    assert first[0] == "SMK-V01" and first[1] == "Expected" and first[2] == "Smoking" and first[3] == "Violation"
    assert first[5] == 30.0 and first[7] == 41.0 and first[8] == "TP"
    # unmatched rows carry the tag in P, no kind
    tagged = [r for r in range(5, 40) if ws.cell(r, 2).value == "Unmatched alert"]
    assert tagged and all(ws.cell(r, 4).value is None for r in tagged)
    assert {"Unscripted – false alarm", "Duplicate", "Late (real act)"} <= {ws.cell(r, 16).value for r in tagged}
    st = wb["Settings"]
    labels = {st.cell(r, 1).value: st.cell(r, 2).value for r in range(1, 30)}
    assert labels["Model version"] == "merged_v2.pt" and labels["Settings consistency"] == "CONSISTENT" \
        or labels["Settings consistency"].startswith("MIXED")
    assert labels["Alert window (seconds)"] == 60 and labels["Count 'Monitoring' as an alert?"] == "No"
    clips = wb["Clips"]
    row = next(r for r in range(5, 60) if clips.cell(r, 1).value == "PRK-V01")
    assert clips.cell(row, 12).value == "Done" and clips.cell(row, 13).value == "Yes" and clips.cell(row, 7).value == "Validation"
