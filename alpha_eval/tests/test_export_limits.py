import sys
import tempfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import db  # noqa: E402
import exporter  # noqa: E402
import fake_data  # noqa: E402


def test_more_new_clips_than_spare_rows_is_an_error():
    conn = db.connect(Path(tempfile.mkdtemp()) / "t.db")
    db.import_template_clips(conn)
    for i in range(96):                                   # rows 60-154 hold 95
        db.upsert_clip(conn, f"NEW-{i:03d}", planned_focus="Smoking", **{"set": "Main"}, in_template_plan=0)
    with pytest.raises(exporter.ExportError) as e:
        exporter.build_export(conn, out_dir=tempfile.mkdtemp(), recalc=False)
    assert "NEW-095" in str(e.value)


def test_results_log_overflow_is_an_error():
    conn = db.connect(Path(tempfile.mkdtemp()) / "t.db")
    db.import_template_clips(conn)
    keys = [("Smoking", "Look-alike", None, "x")] * 401
    fake_data.make_run(conn, "SMK-V01", "Main", [], keys)
    with pytest.raises(exporter.ExportError) as e:
        exporter.build_export(conn, out_dir=tempfile.mkdtemp(), recalc=False)
    assert "401" in str(e.value)


def test_new_clip_lands_in_row_60_and_template_clip_keeps_its_row():
    import openpyxl
    conn = db.connect(Path(tempfile.mkdtemp()) / "t.db")
    db.import_template_clips(conn)
    db.upsert_clip(conn, "SMK-X01", planned_focus="Smoking", scenario="extra", distance="Far", time_of_day="Night",
                   **{"set": "Validation"}, in_template_plan=0, min_length="1 min")
    res = exporter.build_export(conn, out_dir=tempfile.mkdtemp(), recalc=False)
    ws = openpyxl.load_workbook(res["path"])["Clips"]
    assert ws["A60"].value == "SMK-X01" and ws["B60"].value == "Smoking" and ws["G60"].value == "Validation"
    assert ws["A5"].value == "SMK-V01" and ws["L5"].value == "Not run"
