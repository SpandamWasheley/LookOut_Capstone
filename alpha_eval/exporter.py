"""Fill a COPY of the Excel template with the alpha-test results (values only; formulas are never touched).

The template itself is never modified. After writing, the copy is recalculated (LibreOffice when available,
otherwise the `formulas` package) and Results Log column Z ("Check") is read back: any MISMATCH means the tool's
score and the sheet's own formulas disagree, and the export is labelled NOT FINAL.
"""
import json
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

import config
import db
import matching
import scoring

# Columns of the Results Log that are INPUTS (never I, J, N, Q, Z which are formulas)
LOG_INPUT_COLS = [1, 2, 3, 4, 5, 6, 7, 8, 11, 12, 13, 15, 16, 18, 19, 20, 21, 22, 23, 24, 25, 27]
CLIP_FORMULA_COLS = {14}
RUN_STATUS_VALUES = {"queued": "Queued", "running": "Running", "done": "Done", "failed": "Failed", "cancelled": "Not run"}

# Snapshot fields that legitimately differ between Main and Validation (the two test thresholds)
SET_DIFF_OK = {"overrides_in_memory.parking_threshold_sec", "overrides_in_memory.drinking_group_duration_sec",
               "system_settings_used.drinking_group_duration"}
IGNORE_SNAPSHOT = ("worker_inputs.", "video.", "clock_start", "video_clock_patched", "alert_store_patched",
                   "system_settings_used.updated_at", "blocked_writes_active")
SETTINGS_LISTED = {"smoking_dwell", "id"}


class ExportError(Exception):
    pass


def sha256_file(path):
    import hashlib
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for b in iter(lambda: fh.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def r1(v):
    return None if v is None else round(float(v), 1)


def _flatten(d, prefix=""):
    out = {}
    if isinstance(d, dict):
        for k, v in d.items():
            out.update(_flatten(v, f"{prefix}{k}."))
    elif isinstance(d, list):
        out[prefix.rstrip(".")] = json.dumps(d, default=str)
    else:
        out[prefix.rstrip(".")] = d
    return out


def settings_consistency(comps, es):
    """'CONSISTENT' or 'MIXED — <fields that differ>' (compared within each Set, then across Sets)."""
    diffs = set()
    by_set = {}
    for c in comps:
        snap = json.loads(c["run"]["snapshot_json"] or "{}")
        flat = {k: v for k, v in _flatten(snap).items() if not k.startswith(IGNORE_SNAPSHOT)}
        by_set.setdefault(c["clip"]["set"], []).append((c["clip"]["clip_id"], flat))
    for _, items in by_set.items():
        base = items[0][1]
        for _, flat in items[1:]:
            for k in set(base) | set(flat):
                if base.get(k) != flat.get(k):
                    diffs.add(k)
    if len(by_set) == 2:
        a, b = (by_set[s][0][1] for s in by_set)
        for k in set(a) | set(b):
            if a.get(k) != b.get(k) and k not in SET_DIFF_OK:
                diffs.add(k)
    # thresholds actually used vs the thresholds in the Settings sheet
    for c in comps:
        suffix = "validation" if c["clip"]["set"] == "Validation" else "main"
        if float(c["run"]["parking_threshold"] or -1) != float(es[f"parking_threshold_{suffix}_sec"]):
            diffs.add(f"parking_threshold_{suffix}_sec (re-run {c['clip']['clip_id']})")
        if float(c["run"]["drinking_duration"] or -1) != float(es[f"drinking_group_duration_{suffix}_sec"]):
            diffs.add(f"drinking_group_duration_{suffix}_sec (re-run {c['clip']['clip_id']})")
    return "CONSISTENT" if not diffs else "MIXED — " + ", ".join(sorted(diffs))


def _other_settings(snap):
    used = dict(snap.get("system_settings_used") or {})
    for k in ("id", "updated_at", "smoking_dwell"):
        used.pop(k, None)
    return json.dumps(used, ensure_ascii=False, separators=(",", ":"), default=str)


def find_soffice():
    for cand in (shutil.which("soffice"), shutil.which("soffice.exe"),
                 r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if cand and Path(cand).exists():
            return cand
    return None


def recalculate(xlsx_path):
    """Recalculate a copy of the workbook. Returns (values_workbook, method). values_workbook is data_only."""
    soffice = find_soffice()
    tmp = Path(tempfile.mkdtemp(prefix="alpha_recalc_"))
    if soffice:
        src = tmp / "in.xlsx"
        shutil.copyfile(xlsx_path, src)
        out = tmp / "out"
        out.mkdir()
        subprocess.run([soffice, "--headless", "--norestore", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                        "--outdir", str(out), str(src)], capture_output=True, timeout=300)
        res = out / "in.xlsx"
        if res.exists():
            return openpyxl.load_workbook(res, data_only=True), "LibreOffice"
    try:
        import formulas                                   # pure-Python fallback
        xl = formulas.ExcelModel().loads(str(xlsx_path)).finish()
        xl.calculate()
        out = tmp / "formulas_out"
        out.mkdir()
        xl.write(dirpath=str(out))
        produced = next(out.glob("*.XLSX"), None) or next(out.glob("*.xlsx"), None)
        if produced is None:
            raise RuntimeError("formulas wrote no file")
        return openpyxl.load_workbook(produced, data_only=True), "formulas (Python)"
    except Exception as exc:                              # noqa: BLE001
        raise ExportError(f"no recalculation engine available (LibreOffice missing and `formulas` failed: {exc})")


def sheet(wbv, name):
    """Worksheet by name, ignoring case (the `formulas` fallback upper-cases sheet names)."""
    for ws in wbv.worksheets:
        if ws.title.upper() == name.upper():
            return ws
    raise KeyError(name)


def _num(v):
    return v if isinstance(v, (int, float)) else None


def read_sheet_metrics(wbv, sheet_name):
    """Summary / Validation values -> {type: {field: value}} (None for '–')."""
    ws = sheet(wbv, sheet_name)
    fields = ["real_acts", "tp", "fn", "tn", "fp_look", "fp_actor", "fp_unscripted", "fp_total", "accuracy", "precision",
              "recall", "f1", "avg_delay", "duplicates", "late", "unscripted_real", "needs"]
    out = {}
    for r, name in zip(range(5, 10), scoring.TYPES + ["OVERALL"]):
        out[name] = {f: _num(ws.cell(r, 2 + i).value) for i, f in enumerate(fields)}
    return out


def compare_with_preview(wbv, records):
    """Differences between the recalculated Summary/Validation sheets and scoring.aggregate (empty = identical)."""
    diffs = []
    for sheet_name, set_name in (("Summary", "Main"), ("Validation", "Validation")):
        agg = scoring.aggregate([r for r in records if r["set"] == set_name])
        sheetv = read_sheet_metrics(wbv, sheet_name)
        for t, row in sheetv.items():
            for f, sv in row.items():
                pv = agg[t][f]
                if (sv is None) != (pv is None) or (sv is not None and abs(float(sv) - float(pv)) > 1e-9):
                    diffs.append(f"{sheet_name}/{t}/{f}: sheet={sv} preview={pv}")
    return diffs


def build_export(conn, reviewed_only=True, out_dir=None, stamp=None, recalc=True):
    out_dir = Path(out_dir or config.EXPORTS_DIR)
    out_dir.mkdir(exist_ok=True)
    tpl_sha_before = sha256_file(config.TEMPLATE)
    es = db.get_eval_settings(conn)
    runs = matching.eligible_runs(conn, reviewed_only)
    comps = [matching.compute_run(conn, r["run_id"], persist=False) for r in runs]
    comp_by_clip = {c["clip"]["clip_id"]: c for c in comps}
    problems = []

    wb = openpyxl.load_workbook(config.TEMPLATE)               # formulas kept (NOT data_only)
    # ---- Clips sheet ----------------------------------------------------------------------------------
    wsC = wb["Clips"]
    planned_rows = {}
    for r in range(config.PLANNED_ROWS[0], config.PLANNED_ROWS[1] + 1):
        if wsC.cell(r, 1).value:
            planned_rows[str(wsC.cell(r, 1).value)] = r
    clips = db.list_clips(conn)
    new_clips = [c for c in clips if c["clip_id"] not in planned_rows]
    spare = list(range(config.SPARE_ROWS[0], config.SPARE_ROWS[1] + 1))
    if len(new_clips) > len(spare):
        over = [c["clip_id"] for c in new_clips[len(spare):]]
        raise ExportError(f"The Clips sheet has only {len(spare)} spare rows (60–154); no room for: {', '.join(over)}")
    row_of = dict(planned_rows)
    for c, r in zip(new_clips, spare):
        row_of[c["clip_id"]] = r
    for c in clips:
        r = row_of[c["clip_id"]]
        status, latest, done = _clip_status(conn, c)
        comp = comp_by_clip.get(c["clip_id"])
        if c["clip_id"] not in planned_rows:
            wsC.cell(r, 1, c["clip_id"])
            wsC.cell(r, 2, c["planned_focus"])
            wsC.cell(r, 6, c.get("min_length"))
        wsC.cell(r, 3, c["scenario"])
        wsC.cell(r, 4, c["distance"])
        wsC.cell(r, 5, c["time_of_day"])
        wsC.cell(r, 7, c["set"])
        wsC.cell(r, 8, "Recorded" if c["video_path"] else (c["recording_status"] or "Not recorded"))
        wsC.cell(r, 9, Path(c["video_path"]).name if c["video_path"] else None)
        wsC.cell(r, 10, c["date_recorded"])
        wsC.cell(r, 11, c["duration_sec"])
        wsC.cell(r, 12, status)
        wsC.cell(r, 13, "Yes" if (done and done["reviewed"]) else "No")
        if comp:
            run = comp["run"]
            mc = comp["settings"].monitoring_counts_as_alert
            wsC.cell(r, 15, sum(1 for e in comp["events"] if e.counts(mc)))
            wsC.cell(r, 16, None if run["processing_fps"] is None else round(run["processing_fps"], 2))
            wsC.cell(r, 17, None if run["processing_time_sec"] is None else round(run["processing_time_sec"], 1))
            wsC.cell(r, 18, run["run_id"])
        wsC.cell(r, 19, c["notes"])

    # ---- Results Log ----------------------------------------------------------------------------------
    wsL = wb["Results Log"]
    for r in range(config.LOG_ROWS[0], config.LOG_ROWS[1] + 1):
        for col in LOG_INPUT_COLS:
            wsL.cell(r, col).value = None
    records = []
    ordered = sorted(comps, key=lambda c: [x["clip_id"] for x in clips].index(c["clip"]["clip_id"]))
    for c in ordered:
        records.extend(matching.flat_records(c))
    capacity = config.LOG_ROWS[1] - config.LOG_ROWS[0] + 1
    if len(records) > capacity:
        raise ExportError(f"The Results Log needs {len(records)} rows but the template has {capacity} (rows 5–404).")
    reviewer_of = {c["run"]["run_id"]: (c["run"]["reviewer"] if c["run"]["reviewed"] else None) for c in ordered}
    for i, rec in enumerate(records):
        r = config.LOG_ROWS[0] + i
        exp = rec["row_type"] == "Expected"
        vals = {1: rec["clip_id"], 2: rec["row_type"], 3: rec["type"], 4: rec["kind"] if exp else None,
                5: rec["instance_no"] if exp else None, 6: rec["who"], 7: rec["set"],
                8: r1(rec["start"]) if exp else None, 11: rec["event_id"], 12: r1(rec["alert_time"]),
                13: rec["status"] if rec["event_id"] else None, 15: rec["ai"] if rec["event_id"] else None,
                16: rec["tag"] if not exp else None, 18: rec["alerted_instead"] if exp else None,
                19: rec["distance"], 20: rec["time_of_day"],
                21: ("Yes" if rec["added_after_run"] else "No") if exp else None, 22: rec["miss_reason"],
                23: reviewer_of.get(rec["run_id"]), 24: rec["notes"], 25: rec["outcome"], 27: rec["run_id"]}
        for col, v in vals.items():
            wsL.cell(r, col).value = v

    # ---- Settings sheet -------------------------------------------------------------------------------
    wsS = wb["Settings"]
    snaps = [json.loads(c["run"]["snapshot_json"] or "{}") for c in ordered]
    snap0 = snaps[0] if snaps else {}
    consistency = settings_consistency(ordered, es)
    n_needs = sum(1 for r in records if r["outcome"] in ("Needs tag", "Needs start"))
    basis = ("Reviewed clips only" if reviewed_only else "All finished clips (reviewed or not)") + ", latest run per clip"
    git = snap0.get("lookout_git", {})
    values = {
        "Alert window (seconds)": es["alert_window_sec"], "Start tolerance (seconds)": es["start_tolerance_sec"],
        "Count 'Monitoring' as an alert?": "Yes" if es["monitoring_counts_as_alert"] else "No",
        "Parking threshold — Main set (s)": es["parking_threshold_main_sec"],
        "Parking threshold — Validation set (s)": es["parking_threshold_validation_sec"],
        "Drinking group duration — Main set (s)": es["drinking_group_duration_main_sec"],
        "Drinking group duration — Validation set (s)": es["drinking_group_duration_validation_sec"],
        "Model version": snap0.get("model_file"), "Model SHA-256": snap0.get("model_sha256"),
        "Scoring spec": snap0.get("scoring_spec", config.SCORING_SPEC),
        "Smoking dwell (s)": (snap0.get("system_settings_used") or {}).get("smoking_dwell"),
        "Other detection settings": _other_settings(snap0) if snap0 else None,
        "LookOut git commit": (git.get("commit", "") + (" (+dirty)" if git.get("dirty") else "")) if git else None,
        "alpha_eval tool version": config.TOOL_VERSION,
        "Export date": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "Results basis": basis, "Settings consistency": consistency,
    }
    found = set()
    for r in range(1, wsS.max_row + 1):
        label = wsS.cell(r, 1).value
        if label in values:
            if values[label] is not None:
                wsS.cell(r, 2).value = values[label]
            found.add(label)
    missing = set(values) - found
    if missing:
        problems.append("Settings labels not found in the template: " + ", ".join(sorted(missing)))

    # ---- save, recalculate, check ----------------------------------------------------------------------
    stamp = stamp or datetime.now().strftime("%Y%m%d_%H%M")
    wb.calculation.fullCalcOnLoad = True
    path = out_dir / f"LookOut_Alpha_Results_{stamp}.xlsx"
    wb.save(path)
    result = dict(path=str(path), n_clips=len(ordered), n_rows=len(records), mismatches=[], preview_diffs=[],
                  method=None, problems=problems, consistency=consistency, needs=n_needs, records=records)
    if recalc:
        try:
            wbv, method = recalculate(path)
            result["method"] = method
            wsZ = sheet(wbv, "Results Log")
            for i in range(len(records)):
                r = config.LOG_ROWS[0] + i
                if wsZ.cell(r, 26).value == "MISMATCH":
                    rec = records[i]
                    result["mismatches"].append(
                        f"row {r}: {rec['clip_id']} {rec['row_type']} {rec['type']} — tool says {rec['outcome']}, "
                        f"sheet says {wsZ.cell(r, 17).value}")
            result["preview_diffs"] = compare_with_preview(wbv, records)
            result["sheet_values"] = {s: read_sheet_metrics(wbv, s) for s in ("Summary", "Validation")}
        except ExportError as exc:
            problems.append(str(exc))
    else:
        problems.append("recalculation skipped")
    final = (not result["mismatches"]) and (not result["preview_diffs"]) and result["method"] is not None \
        and consistency == "CONSISTENT" and n_needs == 0
    result["final"] = final
    if not final:
        new_path = path.with_name(path.stem + "_NOT_FINAL.xlsx")
        wb2 = openpyxl.load_workbook(path)
        why = []
        if result["mismatches"]:
            why.append(f"{len(result['mismatches'])} MISMATCH")
        if result["preview_diffs"]:
            why.append("sheet/preview differ")
        if consistency != "CONSISTENT":
            why.append("settings not consistent")
        if n_needs:
            why.append(f"{n_needs} row(s) need a tag/start")
        if result["method"] is None:
            why.append("not recalculated")
        ws2 = wb2["Settings"]
        for r in range(1, ws2.max_row + 1):
            if ws2.cell(r, 1).value == "Results basis":
                ws2.cell(r, 2).value = f"NOT FINAL ({'; '.join(why)}) — {basis}"
        wb2.calculation.fullCalcOnLoad = True
        wb2.save(new_path)
        path.unlink()
        result["path"] = str(new_path)
        result["why_not_final"] = why
    if sha256_file(config.TEMPLATE) != tpl_sha_before:
        raise ExportError("the template changed during export — this must never happen")
    result["template_sha256"] = tpl_sha_before
    conn.execute("INSERT INTO exports (created_at, path, final, mismatches, basis, n_clips, note) VALUES (?,?,?,?,?,?,?)",
                 (db.now_iso(), result["path"], int(final), len(result["mismatches"]), basis, len(ordered),
                  "; ".join(result.get("why_not_final", []) + problems)))
    conn.commit()
    return result


def _clip_status(conn, c):
    latest = db.latest_run(conn, c["clip_id"])
    done = db.latest_run(conn, c["clip_id"], "done")
    if latest is None:
        return "Not run", None, None
    status = RUN_STATUS_VALUES.get(latest["status"], "Not run")
    if done and status == "Done":
        keys = db.key_rows(conn, c["clip_id"])
        if any(k["kind"] == "Violation" and k["action_start_sec"] is None for k in keys):
            status = "Needs action start"
    return status, latest, done
